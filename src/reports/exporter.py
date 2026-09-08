from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = PROJECT_ROOT / "output" / "reports"
REPORT_DIR.mkdir(parents=True, exist_ok=True)


def _safe_filename(value: str) -> str:
    """Convert a company/preset name into a safe filename."""
    return "".join(
        character if character.isalnum() or character in "_-" else "_"
        for character in str(value)
    )


def export_excel(
    df: pd.DataFrame,
    report_name: str,
    sheet_name: str = "Report",
) -> Path:
    """Export a DataFrame to a formatted Excel workbook."""

    filename = f"{_safe_filename(report_name)}.xlsx"
    output_path = REPORT_DIR / filename

    export_df = df.copy()

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
    ) as writer:
        export_df.to_excel(
            writer,
            index=False,
            sheet_name=sheet_name[:31],
        )

    workbook = load_workbook(output_path)
    worksheet = workbook[sheet_name[:31]]

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12), 40
        )

    workbook.save(output_path)

    return output_path


def export_pdf(
    df: pd.DataFrame,
    report_name: str,
    title: str = "N100 Financial Intelligence Report",
) -> Path:
    """Export a DataFrame to a landscape PDF report."""

    filename = f"{_safe_filename(report_name)}.pdf"
    output_path = REPORT_DIR / filename

    document = SimpleDocTemplate(
        str(output_path),
        pagesize=landscape(A4),
        rightMargin=24,
        leftMargin=24,
        topMargin=24,
        bottomMargin=24,
    )

    styles = getSampleStyleSheet()

    story = []

    story.append(
        Paragraph(
            title,
            styles["Title"],
        )
    )

    story.append(
        Paragraph(
            f"Generated: {datetime.now():%Y-%m-%d %H:%M:%S}",
            styles["Normal"],
        )
    )

    story.append(Spacer(1, 12))

    display_df = df.copy()

    # Limit very wide reports for PDF readability.
    if len(display_df.columns) > 10:
        display_df = display_df.iloc[:, :10]

    display_df = display_df.fillna("")

    table_data = [[str(column) for column in display_df.columns]]

    for row in display_df.astype(str).values.tolist():
        table_data.append(row)

    table = Table(
        table_data,
        repeatRows=1,
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#1F4E78"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#F2F2F2"),
                    ],
                ),
            ]
        )
    )

    story.append(table)

    document.build(story)

    return output_path
